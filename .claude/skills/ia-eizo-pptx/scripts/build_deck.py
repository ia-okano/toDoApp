#!/usr/bin/env python3
"""IA 映像教材フォーマット (デザインシステム) に沿った pptx を生成するツール。

元の "FMT" テンプレートのスライドマスター / レイアウト / テーマ / 画像
(base_template.pptx) をそのまま再利用し、その上に新しいスライドを構築する。
これによりデザイン (扉ページ・上部バー・配色・フォント) が完全に踏襲される。

使い方:
    python build_deck.py <spec.json> [-o output.pptx] [--base base_template.pptx]

spec.json の形式 (詳細は SKILL.md / examples を参照):
{
  "output": "deck.pptx",            # 省略可 (-o で上書き可)
  "slides": [
    {"type": "cover",   "title": "プレゼンテーション タイトル"},
    {"type": "section", "title": "第1章 セクションタイトル"},
    {"type": "content", "title": "ページ見出し(L1)",
        "heading2": "セクション見出し(L2)",
        "heading3": "小見出し(L3)",
        "body": [
            {"text": "本文テキスト。"},
            {"text": "箇条書き項目", "bullet": true},
            {"text": "第2階層", "bullet": true, "level": 1, "size": 2000},
            {"runs": [
                {"text": "強調は"}, {"text": "太字", "bold": true},
                {"text": "、重要は"}, {"text": "赤太字", "bold": true, "color": "FF4B4B"},
                {"text": "を使用。"}
            ]}
        ]},
    {"type": "two_column", "title": "2カラム",
        "left":  {"heading3": "左見出し", "body": [{"text": "左本文"}]},
        "right": {"heading3": "右見出し", "body": [{"text": "右本文"}]}}
  ]
}

色は schemeClr 名 (bg1, tx1, tx2, accent1..6) か 6桁 16進 (例 "FF4B4B")。
デザインシステムの既定値:
  L1 (上部バー)  : 28pt 太字 白(bg1)
  L2 (見出し)    : 28pt 太字 シアン(tx2 = #2BACED)
  L3 (小見出し)  : 24pt 太字 ダーク(tx1 = #333333)
  本文           : 24pt 細字 ダーク(tx1)、行間 1.3、最小 16pt 目安
  ページ番号     : 右下 (マスター既定)、表紙・中表紙には付けない
"""
import argparse, json, os, sys
from pptx import Presentation
from pptx.oxml.ns import qn, nsdecls
from pptx.oxml import parse_xml

YU = ('<a:latin typeface="游ゴシック" panose="020B0400000000000000" pitchFamily="50" charset="-128"/>'
      '<a:ea typeface="游ゴシック" panose="020B0400000000000000" pitchFamily="50" charset="-128"/>')

# ---- design system defaults (EMU / 1pt = 100) ----
BAR_H        = 783771      # 上部バー高さ (~2.18cm)
MARGIN_X     = 487363      # 左マージン (~1.35cm)
COL_W        = 10944225    # 1カラム本文幅 (~30.4cm)
WIDE_W       = 11217600    # 広めの本文幅 (~31.16cm)
SLIDE_W      = 12192000
L2_Y         = 1196975     # 見出しL2 の y (~3.32cm)
SZ_L1 = SZ_L2 = 2800
SZ_L3 = SZ_BODY = 2400

_next_id = [1000]
def _nid():
    _next_id[0] += 1
    return _next_id[0]

def _esc(s):
    return (s or '').replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')

def _fill(color):
    if not color:
        color = 'tx1'
    scheme = {'bg1','bg2','tx1','tx2','dk1','dk2','lt1','lt2',
              'accent1','accent2','accent3','accent4','accent5','accent6'}
    if color in scheme:
        return f'<a:solidFill><a:schemeClr val="{color}"/></a:solidFill>'
    return f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'

def _run(text, size=SZ_BODY, bold=False, color='tx1'):
    b = ' b="1"' if bold else ''
    return (f'<a:r><a:rPr kumimoji="1" lang="ja-JP" altLang="en-US" sz="{size}"{b} dirty="0">'
            f'{_fill(color)}{YU}</a:rPr><a:t>{_esc(text)}</a:t></a:r>')

def _runs_from(item, default_size, default_color='tx1'):
    """item は {"text":...} か {"runs":[...]}。run XML を返す。"""
    if 'runs' in item:
        out = ''
        for r in item['runs']:
            out += _run(r['text'], r.get('size', default_size),
                        r.get('bold', False), r.get('color', default_color))
        return out
    return _run(item['text'], item.get('size', default_size),
                item.get('bold', False), item.get('color', default_color))

def _para(item, default_size):
    lnspc = item.get('linespc', 130000)
    lvl = item.get('level', 0)
    if item.get('bullet'):
        marL = 228600 + lvl * 457200
        ppr = (f'<a:pPr marL="{marL}" indent="-228600" lvl="{lvl}">'
               f'<a:lnSpc><a:spcPct val="{lnspc}"/></a:lnSpc>'
               f'<a:buFont typeface="Arial"/><a:buChar char="•"/></a:pPr>')
    else:
        ppr = (f'<a:pPr lvl="{lvl}">'
               f'<a:lnSpc><a:spcPct val="{lnspc}"/></a:lnSpc><a:buNone/></a:pPr>')
    return f'<a:p>{ppr}{_runs_from(item, item.get("size", default_size))}</a:p>'

def _append(slide, xml):
    slide.shapes._spTree.append(parse_xml(xml))

def _title_bar(slide, text):
    runs = _run(text, SZ_L1, True, 'bg1')
    xml = (f'<p:sp {nsdecls("p","a")}><p:nvSpPr><p:cNvPr id="{_nid()}" name="見出しレベル1"/>'
           f'<p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm><a:off x="0" y="0"/>'
           f'<a:ext cx="{SLIDE_W}" cy="{BAR_H}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
           f'<a:noFill/><a:ln><a:noFill/></a:ln></p:spPr><p:txBody>'
           f'<a:bodyPr lIns="288000" tIns="72000" bIns="36000" rtlCol="0" anchor="ctr"/>'
           f'<a:lstStyle/><a:p><a:pPr lvl="0"/>{runs}</a:p></p:txBody></p:sp>')
    _append(slide, xml)

def _textbox(slide, name, x, y, cx, cy, paras_xml, autofit=False):
    af = '<a:spAutoFit/>' if autofit else ''
    xml = (f'<p:sp {nsdecls("p","a")}><p:nvSpPr><p:cNvPr id="{_nid()}" name="{_esc(name)}"/>'
           f'<p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm><a:off x="{x}" y="{y}"/>'
           f'<a:ext cx="{cx}" cy="{cy}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
           f'<a:noFill/></p:spPr><p:txBody><a:bodyPr wrap="square" rtlCol="0">{af}</a:bodyPr>'
           f'<a:lstStyle/>{paras_xml}</p:txBody></p:sp>')
    _append(slide, xml)

def _add_sldnum(slide):
    xml = (f'<p:sp {nsdecls("p","a")}><p:nvSpPr><p:cNvPr id="{_nid()}" name="スライド番号プレースホルダー"/>'
           f'<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
           f'<p:nvPr><p:ph type="sldNum" sz="quarter" idx="10"/></p:nvPr></p:nvSpPr>'
           f'<p:spPr/><p:txBody><a:bodyPr/><a:lstStyle/><a:p>'
           f'<a:fld id="{{F599A5BA-B7C2-4E23-AF82-85FD31EF0D4A}}" type="slidenum">'
           f'<a:rPr lang="ja-JP" altLang="en-US" smtClean="0"/><a:pPr/><a:t>1</a:t></a:fld>'
           f'<a:endParaRPr lang="ja-JP" altLang="en-US"/></a:p></p:txBody></p:sp>')
    _append(slide, xml)

def _set_title(slide, text):
    for ph in slide.placeholders:
        pf = ph.placeholder_format
        if pf.idx == 0 or (pf.type is not None and 'TITLE' in str(pf.type)):
            ph.text = text
            return
    phs = list(slide.placeholders)
    if phs:
        phs[0].text = text

def _body_block(slide, name, x, y, cx, cy, body_items):
    paras = ''.join(_para(it, SZ_BODY) for it in body_items)
    _textbox(slide, name, x, y, cx, cy, paras, autofit=False)

def build(spec, base_path, out_path):
    prs = Presentation(base_path)
    layouts = {l.name: l for l in prs.slide_masters[0].slide_layouts}

    for sl in spec['slides']:
        t = sl['type']
        if t == 'cover':
            s = prs.slides.add_slide(layouts['表紙'])
            _set_title(s, sl.get('title', ''))
        elif t == 'section':
            s = prs.slides.add_slide(layouts['中表紙'])
            _set_title(s, sl.get('title', ''))
        elif t == 'content':
            s = prs.slides.add_slide(layouts['1カラム'])
            _title_bar(s, sl.get('title', ''))
            y = L2_Y
            if sl.get('heading2'):
                _textbox(s, '見出しレベル2', MARGIN_X, y, COL_W, 523220,
                         f'<a:p><a:pPr lvl="0"/>{_run(sl["heading2"], SZ_L2, True, "tx2")}</a:p>')
                y += 643000
            if sl.get('heading3'):
                _textbox(s, '見出しレベル3', MARGIN_X, y, COL_W, 430000,
                         f'<a:p><a:pPr lvl="0"/>{_run(sl["heading3"], SZ_L3, True, "tx1")}</a:p>')
                y += 560000
            if sl.get('body'):
                _body_block(s, '本文', MARGIN_X + 5837, y, WIDE_W, 6858000 - y - 600000, sl['body'])
            _add_sldnum(s)
        elif t == 'two_column':
            s = prs.slides.add_slide(layouts['２カラム'])
            _title_bar(s, sl.get('title', ''))
            half = 5400000
            lx, rx = MARGIN_X, 6300000
            for col, x in (('left', lx), ('right', rx)):
                c = sl.get(col, {})
                yy = 1300000
                if c.get('heading3'):
                    _textbox(s, f'見出しレベル3_{col}', x, yy, half, 430000,
                             f'<a:p><a:pPr lvl="0"/>{_run(c["heading3"], SZ_L3, True, "tx1")}</a:p>')
                    yy += 600000
                if c.get('body'):
                    _body_block(s, f'本文_{col}', x, yy, half, 6858000 - yy - 600000, c['body'])
            _add_sldnum(s)
        else:
            raise ValueError(f'unknown slide type: {t}')

    prs.save(out_path)
    return len(prs.slides._sldIdLst)

def main():
    ap = argparse.ArgumentParser(description='IA 映像教材フォーマットの pptx を生成')
    ap.add_argument('spec', help='スライド定義 JSON')
    ap.add_argument('-o', '--output', help='出力 pptx パス')
    ap.add_argument('--base', help='base_template.pptx のパス')
    args = ap.parse_args()

    here = os.path.dirname(os.path.abspath(__file__))
    base = args.base or os.path.join(here, '..', 'assets', 'base_template.pptx')
    with open(args.spec, encoding='utf-8') as f:
        spec = json.load(f)
    out = args.output or spec.get('output', 'deck.pptx')
    n = build(spec, base, out)
    print(f'OK: {out} ({n} slides)')

if __name__ == '__main__':
    main()
